"""VIC suburb crime — Crime Statistics Agency Victoria, Recorded Offences "Table 03".

The VIC third of the crime family (crime = three state feeds: BOCSAR NSW [built],
CSA VIC [this], QPS QLD [deferred]). Same metric MECHANISM as bocsar — sum all recorded
offences over the latest annual window per suburb → /census_total_persons * 1000 =
crime_incidents_per_1000. Map-only raw, NO resolver band ([[adapter-band-sourced-vs-projected]]
case-3: a band would stamp "high crime" on the commercial hubs).

FETCH is **human-staged**, not a live pull: CSA's HTML page is UA-filtered (Drupal, not a
Cloudflare wall — the file host files.crimestatistics.vic.gov.au is reachable, so this is
live-pullable later), but quarterly cadence makes a manual drop legitimate, and the staged
.xlsx is COMMITTED for reproducibility (the vic_vpsr discipline). Staged file:

    engine/build/suburbs/data/Data_Tables_LGA_Recorded_Offences_Year_Ending_*.xlsx

The suburb grain lives in sheet 'Table 03' (Offences recorded by offence type, LGA and
postcode or suburb/town). Rows are annual (year ending December), one per
(year, LGA, postcode, suburb, offence subgroup); we keep the latest year and sum Offence
Count across postcodes + subgroups per (suburb, LGA), then resolve suburb→SAL via the shared
name→SAL crosswalk (_namecross; LGA disambiguates name duplicates — CSA's clean LGA column
beats VPSR's paren-qualifier). .xlsx → openpyxl (contrast vic_vpsr's legacy .xls/xlrd).

Run:  cd engine/build && ENGINE_DATABASE_URL=… ../../.venv/bin/python -m suburbs.csa_vic
"""
from __future__ import annotations

import glob
import os
import sys
from collections import defaultdict

import openpyxl

from . import _namecross, db

DATA_DIR = os.path.join(os.path.dirname(__file__), "data")
FILE_GLOB = "Data_Tables_LGA_Recorded_Offences_*.xlsx"
SHEET = "Table 03"

SOURCE = {
    "source_id": "csa_vic",
    "name": "Recorded Offences by suburb/town – LGA Recorded Offences, Table 03",
    "publisher": "Crime Statistics Agency Victoria",
    "license": "CC BY 4.0",
    "attribution": "© State of Victoria (Crime Statistics Agency)",
    "redistribution": "permitted",
    "cadence": "quarterly",
    "url": "https://www.crimestatistics.vic.gov.au/crime-statistics/latest-victorian-crime-data/download-data",
    "notes": "Human-staged .xlsx (Table 03, suburb/town grain). Annual offence counts (year ending "
             "December) summed across subgroups + postcodes per suburb → /census_total_persons * 1000 = "
             "crime_incidents_per_1000 (VIC). Map-only raw, no resolver band (§1). NSW=bocsar, QLD deferred.",
}


def _staged_file() -> str:
    hits = sorted(glob.glob(os.path.join(DATA_DIR, FILE_GLOB)))
    if not hits:
        raise SystemExit(f"csa_vic: no staged file matching {FILE_GLOB} in {DATA_DIR} "
                         "(download the CSA 'LGA Recorded Offences' .xlsx there)")
    return hits[-1]  # latest by name — the reporting period is in the filename


def _parse(path: str) -> tuple[dict[tuple[str, str], int], str, str]:
    """Stream 'Table 03' → ({(suburb, lga): annual_offence_count}, as_of, period).

    Keeps only the latest 'year ending December' window and sums Offence Count across
    all offence subgroups and postcodes per (suburb, LGA)."""
    wb = openpyxl.load_workbook(path, read_only=True)
    if SHEET not in wb.sheetnames:
        wb.close()
        raise SystemExit(f"csa_vic: sheet {SHEET!r} not in {path} — layout changed?")
    ws = wb[SHEET]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(c).strip() if c is not None else "" for c in next(rows)]
    try:
        i_yr = hdr.index("Year")
        i_ye = hdr.index("Year ending")
        i_lga = hdr.index("Local Government Area")
        i_sub = hdr.index("Suburb/Town Name")
        i_cnt = hdr.index("Offence Count")
    except ValueError as e:
        wb.close()
        raise SystemExit(f"csa_vic: Table 03 header changed ({e}); got {hdr}")

    by_year: dict[int, dict[tuple[str, str], int]] = defaultdict(lambda: defaultdict(int))
    for r in rows:
        if r[i_ye] != "December":
            continue
        try:
            year = int(r[i_yr])
        except (TypeError, ValueError):
            continue
        sub = str(r[i_sub]).strip()
        if not sub:
            continue
        lga = str(r[i_lga]).strip()
        try:
            cnt = int(r[i_cnt])
        except (TypeError, ValueError):
            cnt = 0  # CSA masks sensitive subgroups (<3) → non-numeric; treat as 0 (honest, can't recover)
        by_year[year][(sub, lga)] += cnt
    wb.close()
    if not by_year:
        raise SystemExit("csa_vic: no 'December' rows in Table 03 — layout changed?")
    latest = max(by_year)
    return dict(by_year[latest]), f"{latest}-12-31", f"January {latest} – December {latest}"


def _population(conn) -> dict[str, int]:
    """{sal_code: census_total_persons} for VIC — the rate denominator (§8)."""
    pop: dict[str, int] = {}
    with conn.cursor() as cur:
        cur.execute(
            "SELECT sal_code, (facts_jsonb->>'census_total_persons')::int "
            "FROM suburbs WHERE state = 'VIC' AND facts_jsonb ? 'census_total_persons'"
        )
        for sal, persons in cur.fetchall():
            pop[sal] = persons
    return pop


def run() -> tuple[int, int]:
    agg, as_of, period = _parse(_staged_file())
    conn = db.connect()
    try:
        xw = _namecross.build(conn, "VIC", "Vic.")
        pop = _population(conn)
        src = {"source_id": SOURCE["source_id"], "as_of": as_of}

        # Resolve (suburb, LGA) → SAL and accumulate by SAL: a suburb spanning LGAs/postcodes
        # folds into one SAL; genuine name-duplicates resolve to distinct SALs.
        by_sal: dict[str, int] = defaultdict(int)
        dropped: list[str] = []
        for (sub, lga), n in agg.items():
            sal = xw.resolve(f"{sub} ({lga})")
            if sal is None:
                dropped.append(f"{sub} ({lga})")
                continue
            by_sal[sal] += n

        rows: list[dict] = []
        no_pop = 0
        for sal, annual in by_sal.items():
            facts = {"crime_incidents_annual": annual, "crime_period": period}
            prov = {"crime_incidents_annual": src, "crime_period": src}
            persons = pop.get(sal)
            if persons and persons > 0:
                facts["crime_incidents_per_1000"] = round(annual / persons * 1000, 1)
                prov["crime_incidents_per_1000"] = src
            else:
                no_pop += 1  # honest-partial: count stored, rate null (no denominator)
            rows.append({"sal_code": sal, "facts": facts, "prov": prov})

        db.ensure_source(conn, SOURCE)
        matched, missed = db.update_facts(conn, rows)
        conn.commit()
    finally:
        conn.close()

    print(f"  window {period} (as_of {as_of}); {len(agg)} CSA (suburb,LGA) keys → {len(by_sal)} SALs",
          file=sys.stderr)
    if dropped:
        print(f"  {len(dropped)} CSA localities had no SAL match (e.g. {dropped[:5]})", file=sys.stderr)
    if no_pop:
        print(f"  {no_pop} matched SALs had no population → rate null (count kept)", file=sys.stderr)
    if missed:
        print(f"  {len(missed)} SALs had no spine row (unexpected — e.g. {missed[:5]})", file=sys.stderr)
    return matched, len(missed)


if __name__ == "__main__":
    matched, missed = run()
    print(f"csa_vic: crime rate on {matched} VIC suburbs ({missed} unmatched spine rows)")
